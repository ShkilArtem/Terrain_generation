import csv
import math
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference, ScatterChart, Series
    from openpyxl.chart.label import DataLabelList
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ModuleNotFoundError:
    codex_site_packages = Path.home() / ".cache" / "codex-runtimes" / "codex-primary-runtime" / "dependencies" / "python" / "Lib" / "site-packages"
    if codex_site_packages.exists():
        sys.path.insert(0, str(codex_site_packages))
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, LineChart, PieChart, Reference, ScatterChart, Series
        from openpyxl.chart.label import DataLabelList
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    else:
        raise ModuleNotFoundError(
            "openpyxl is required to generate the Excel workbook. "
            "Install it with: python -m pip install openpyxl"
        )


PROJECT_ROOT = Path(__file__).resolve().parents[1]
BENCHMARK_DIR = PROJECT_ROOT / "benchmark_results"
DEFAULT_CSV_PATH = BENCHMARK_DIR / "terrain_metrics_v2.csv"


def to_number(value):
    try:
        if value == "":
            return None
        return float(value)
    except ValueError:
        return value


def mean(values):
    values = [value for value in values if isinstance(value, (int, float))]
    return sum(values) / len(values) if values else None


def correlation(xs, ys):
    pairs = [(x, y) for x, y in zip(xs, ys) if isinstance(x, (int, float)) and isinstance(y, (int, float))]
    if len(pairs) < 2:
        return None
    xs, ys = zip(*pairs)
    mean_x = sum(xs) / len(xs)
    mean_y = sum(ys) / len(ys)
    numerator = sum((x - mean_x) * (y - mean_y) for x, y in pairs)
    denom_x = math.sqrt(sum((x - mean_x) ** 2 for x in xs))
    denom_y = math.sqrt(sum((y - mean_y) ** 2 for y in ys))
    if denom_x == 0 or denom_y == 0:
        return None
    return numerator / (denom_x * denom_y)


def style_header(ws, row=1):
    fill = PatternFill("solid", fgColor="1F4E79")
    for cell in ws[row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill


def autofit(ws, max_width=34):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = min(max_width, max(10, max(len(str(cell.value or "")) for cell in col) + 2))
        ws.column_dimensions[letter].width = width


def append_table(ws, headers, rows, start_row=1, start_col=1):
    for col, header in enumerate(headers, start=start_col):
        ws.cell(start_row, col, header)
    for row_index, row in enumerate(rows, start=start_row + 1):
        for col_index, value in enumerate(row, start=start_col):
            ws.cell(row_index, col_index, value)
    style_header(ws, start_row)
    return start_row, start_row + len(rows)


def nice_major_unit(values, target_ticks=5):
    numeric_values = [abs(value) for value in values if isinstance(value, (int, float))]
    if not numeric_values:
        return None
    max_value = max(numeric_values)
    if max_value <= 0:
        return None
    raw_step = max_value / max(1, target_ticks)
    magnitude = 10 ** math.floor(math.log10(raw_step))
    for multiplier in (1, 2, 5, 10):
        step = multiplier * magnitude
        if raw_step <= step:
            return step
    return 10 * magnitude


def configure_chart(chart, x_title=None, y_title=None, x_format="0", y_format="0.00",
                    x_major_unit=None, y_major_unit=None, legend_position="b"):
    if x_title:
        chart.x_axis.title = x_title
    if y_title:
        chart.y_axis.title = y_title
    chart.x_axis.numFmt = x_format
    chart.y_axis.numFmt = y_format
    chart.x_axis.majorTickMark = "out"
    chart.y_axis.majorTickMark = "out"
    chart.x_axis.tickLblPos = "nextTo"
    chart.y_axis.tickLblPos = "nextTo"
    chart.x_axis.majorGridlines = None
    if x_major_unit:
        chart.x_axis.majorUnit = x_major_unit
    if y_major_unit:
        chart.y_axis.majorUnit = y_major_unit
    if chart.legend:
        chart.legend.position = legend_position


def add_line_chart(ws, title, data_range, category_range, anchor, x_title=None, y_title=None,
                   x_format="0", y_format="0.00", y_major_unit=None):
    chart = LineChart()
    chart.title = title
    chart.add_data(data_range, titles_from_data=True)
    chart.set_categories(category_range)
    chart.height = 10
    chart.width = 20
    configure_chart(chart, x_title, y_title, x_format=x_format, y_format=y_format, y_major_unit=y_major_unit)
    ws.add_chart(chart, anchor)
    return chart


def add_scatter_chart(ws, title, x_ref, y_ref, anchor, x_title, y_title,
                      x_format="0", y_format="0.00", x_major_unit=None, y_major_unit=None,
                      series_title=None):
    chart = ScatterChart()
    chart.title = title
    chart.series.append(Series(y_ref, x_ref, title=series_title or y_title))
    chart.height = 10
    chart.width = 20
    configure_chart(chart, x_title, y_title, x_format=x_format, y_format=y_format,
                    x_major_unit=x_major_unit, y_major_unit=y_major_unit)
    ws.add_chart(chart, anchor)
    return chart


def main():
    csv_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_CSV_PATH
    if not csv_path.is_absolute():
        csv_path = PROJECT_ROOT / csv_path
    if not csv_path.exists():
        raise FileNotFoundError(f"Missing benchmark CSV: {csv_path}")

    BENCHMARK_DIR.mkdir(exist_ok=True)
    xlsx_path = BENCHMARK_DIR / f"{csv_path.stem}_analysis.xlsx"

    with csv_path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        rows = [{key: to_number(value) for key, value in row.items()} for row in reader]
    if not rows:
        raise ValueError("CSV has a header but no benchmark rows yet.")

    for i, row in enumerate(rows, start=1):
        row.setdefault("Run", i)
        row.setdefault("ExperimentalValue", row.get("Iterations"))
        row.setdefault("ExperimentalParameter", "Benchmark")
        row.setdefault("MemoryEstimate_MB", estimate_memory_mb(row.get("GridSize") or 0))

    headers = list(rows[0].keys())
    bin_headers = [f"Bin{i}" for i in range(90)]
    hypso_headers = [f"HypsoAbove{i}" for i in range(101)]
    derived_headers = [
        "Run",
        "ExperimentalParameter",
        "ExperimentalValue",
        "GridSize",
        "Iterations",
        "GenTime_ms",
        "ErosionTime_ms",
        "ThermalTime_ms",
        "PhysicsTime_ms",
        "TotalTime_ms",
        "GenTime_pct",
        "ErosionTime_pct",
        "ThermalTime_pct",
        "MemoryEstimate_MB",
        "TotalVolumeMoved",
        "MeanSlope_deg",
        "P95Slope_deg",
        "HypsometricAUC",
    ]
    derived_rows = [derive_row(row, bin_headers, hypso_headers) for row in rows]
    latest = rows[-1]
    latest_derived = derived_rows[-1]

    wb = Workbook()
    ws_readme = wb.active
    ws_readme.title = "README"
    ws_raw = wb.create_sheet("Raw Data")
    ws_derived = wb.create_sheet("Derived Metrics")
    ws_grid = wb.create_sheet("Cost Grid Size")
    ws_iter = wb.create_sheet("Cost Iterations")
    ws_profile = wb.create_sheet("Profiling")
    ws_realism = wb.create_sheet("Realism Metrics")
    ws_tradeoff = wb.create_sheet("Tradeoff Elbow")
    ws_corr = wb.create_sheet("Correlations")

    write_readme(ws_readme)

    raw_headers = headers
    append_table(ws_raw, raw_headers, [[row.get(header) for header in raw_headers] for row in rows])
    ws_raw.freeze_panes = "A2"

    append_table(ws_derived, derived_headers, [[row.get(header) for header in derived_headers] for row in derived_rows])
    ws_derived.freeze_panes = "A2"

    write_cost_grid_sheet(ws_grid, derived_rows)
    write_cost_iterations_sheet(ws_iter, derived_rows)
    write_profiling_sheet(ws_profile, latest_derived)
    write_realism_sheet(ws_realism, latest, rows, derived_rows, bin_headers, hypso_headers)
    write_tradeoff_sheet(ws_tradeoff, derived_rows)
    write_correlations_sheet(ws_corr, derived_rows)

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        autofit(ws)

    wb.save(xlsx_path)
    print(f"Saved {xlsx_path}")


def estimate_memory_mb(grid_size):
    try:
        n = int(grid_size)
    except (TypeError, ValueError):
        return None
    vertex_bytes = n * n * 16 * 4
    index_bytes = max(0, n - 1) * max(0, n - 1) * 6 * 4
    working_bytes = n * n * 4
    return (vertex_bytes + index_bytes + working_bytes) / (1024 * 1024)


def derive_row(row, bin_headers, hypso_headers):
    bins = [row.get(header, 0) or 0 for header in bin_headers]
    total_bins = sum(bins) or 1
    mean_slope = sum(count * (degree + 0.5) for degree, count in enumerate(bins)) / total_bins
    cumulative = 0
    p95_slope = 89
    for degree, count in enumerate(bins):
        cumulative += count
        if cumulative / total_bins >= 0.95:
            p95_slope = degree
            break
    hypso_values = [row.get(header, 0) or 0 for header in hypso_headers]
    hypso_auc = sum(hypso_values) / len(hypso_values)
    gen = row.get("GenTime_ms") or 0
    erosion = row.get("ErosionTime_ms") or 0
    thermal = row.get("ThermalTime_ms") or 0
    physics_time = erosion + thermal
    total_time = gen + erosion + thermal
    return {
        "Run": row.get("Run"),
        "ExperimentalParameter": row.get("ExperimentalParameter", "Benchmark"),
        "ExperimentalValue": row.get("ExperimentalValue"),
        "GridSize": row.get("GridSize"),
        "Iterations": row.get("Iterations"),
        "GenTime_ms": gen,
        "ErosionTime_ms": erosion,
        "ThermalTime_ms": thermal,
        "PhysicsTime_ms": physics_time,
        "TotalTime_ms": total_time,
        "GenTime_pct": 100 * gen / total_time if total_time else 0,
        "ErosionTime_pct": 100 * erosion / total_time if total_time else 0,
        "ThermalTime_pct": 100 * thermal / total_time if total_time else 0,
        "MemoryEstimate_MB": row.get("MemoryEstimate_MB") or estimate_memory_mb(row.get("GridSize")),
        "TotalVolumeMoved": row.get("TotalVolumeMoved"),
        "MeanSlope_deg": mean_slope,
        "P95Slope_deg": p95_slope,
        "HypsometricAUC": hypso_auc,
    }


def write_readme(ws):
    rows = [
        ["Sheet", "Purpose"],
        ["Raw Data", "Direct CSV import. One benchmark or sweep step equals one row."],
        ["Derived Metrics", "Adds total runtime (ms), physics runtime (ms), stage percentages (%), mean slope (deg), P95 slope (deg), hypsometric AUC (%), and memory estimate (MB)."],
        ["Cost Grid Size", "Use for Coût de calcul: time (ms) and memory (MB) versus fixed GRID_SIZE (vertices per side)."],
        ["Cost Iterations", "Use for Coût de calcul: runtime (ms) versus droplet count K. Expected hydraulic erosion trend is O(K)."],
        ["Profiling", "Pie chart of generation, hydraulic erosion, and thermal weathering time share (% of total run time)."],
        ["Realism Metrics", "Slope histogram (degrees), hypsometric curve (% area above height), and volume moved trend (height units)."],
        ["Tradeoff Elbow", "Main TIPE compromise graph: physics simulation cost (ms) versus realism gain (moved volume)."],
        ["Correlations", "Pearson correlation matrix between runtime, memory, volume moved, and morphology metrics."],
        [],
        ["How to run", "1. Run benchmarks/sweeps in the app. 2. Run: python tools/generate_benchmark_excel.py benchmark_results/parameter_sweep_report_v2.csv"],
    ]
    for row in rows:
        ws.append(row)
    style_header(ws)
    ws["A1"].font = Font(bold=True, color="FFFFFF")


def grouped_average(rows, key):
    groups = {}
    for row in rows:
        value = row.get(key)
        if value is None:
            continue
        groups.setdefault(value, []).append(row)
    output = []
    for value in sorted(groups):
        group = groups[value]
        output.append([
            value,
            mean([r.get("GenTime_ms") for r in group]),
            mean([r.get("ErosionTime_ms") for r in group]),
            mean([r.get("ThermalTime_ms") for r in group]),
            mean([r.get("TotalTime_ms") for r in group]),
            mean([r.get("MemoryEstimate_MB") for r in group]),
            mean([r.get("TotalVolumeMoved") for r in group]),
        ])
    return output


def write_cost_grid_sheet(ws, rows):
    grid_rows = [r for r in rows if r.get("ExperimentalParameter") == "Grid Size"] or rows
    table = grouped_average(grid_rows, "GridSize")
    headers = ["GRID_SIZE (vertices/side)", "fBm generation (ms)", "Hydraulic erosion (ms)", "Thermal weathering (ms)", "Total runtime (ms)", "Estimated memory (MB)", "Moved volume (height units)"]
    start, end = append_table(ws, headers, table)
    if table:
        time_values = [value for row in table for value in row[1:5]]
        memory_values = [row[5] for row in table]
        add_line_chart(
            ws,
            "Time vs GRID_SIZE",
            Reference(ws, min_col=2, max_col=5, min_row=start, max_row=end),
            Reference(ws, min_col=1, min_row=start + 1, max_row=end),
            "I2",
            x_title="GRID_SIZE (vertices per side)",
            y_title="Execution time (ms)",
            x_format="0",
            y_format="0",
            y_major_unit=nice_major_unit(time_values),
        )
        add_line_chart(
            ws,
            "Estimated Memory vs GRID_SIZE",
            Reference(ws, min_col=6, min_row=start, max_row=end),
            Reference(ws, min_col=1, min_row=start + 1, max_row=end),
            "I20",
            x_title="GRID_SIZE (vertices per side)",
            y_title="Estimated memory (MB)",
            x_format="0",
            y_format="0.0",
            y_major_unit=nice_major_unit(memory_values),
        )


def write_cost_iterations_sheet(ws, rows):
    iteration_rows = [r for r in rows if r.get("ExperimentalParameter") == "Iteration Count"] or rows
    table = grouped_average(iteration_rows, "Iterations")
    headers = ["Droplet iterations (count)", "fBm generation (ms)", "Hydraulic erosion (ms)", "Thermal weathering (ms)", "Total runtime (ms)", "Estimated memory (MB)", "Moved volume (height units)"]
    start, end = append_table(ws, headers, table)
    if table:
        time_values = [value for row in table for value in row[2:5]]
        iteration_values = [row[0] for row in table]
        volume_values = [row[6] for row in table]
        add_line_chart(
            ws,
            "Time vs Droplet Iterations",
            Reference(ws, min_col=3, max_col=5, min_row=start, max_row=end),
            Reference(ws, min_col=1, min_row=start + 1, max_row=end),
            "I2",
            x_title="Droplet iterations (count)",
            y_title="Execution time (ms)",
            x_format="0",
            y_format="0",
            y_major_unit=nice_major_unit(time_values),
        )
        add_scatter_chart(
            ws,
            "Iterations vs Volume Moved",
            Reference(ws, min_col=1, min_row=start + 1, max_row=end),
            Reference(ws, min_col=7, min_row=start + 1, max_row=end),
            "I20",
            "Droplet iterations (count)",
            "Moved volume (height units)",
            x_format="0",
            y_format="0.00",
            x_major_unit=nice_major_unit(iteration_values),
            y_major_unit=nice_major_unit(volume_values),
            series_title="Moved volume (height units)",
        )


def write_profiling_sheet(ws, latest):
    rows = [
        ["Stage", "Time (ms)", "Share of total time (%)"],
        ["fBm noise generation", latest.get("GenTime_ms"), latest.get("GenTime_pct")],
        ["Hydraulic erosion", latest.get("ErosionTime_ms"), latest.get("ErosionTime_pct")],
        ["Thermal weathering", latest.get("ThermalTime_ms"), latest.get("ThermalTime_pct")],
    ]
    for row in rows:
        ws.append(row)
    style_header(ws)
    pie = PieChart()
    pie.title = "Profiling: Time Share by Stage"
    pie.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
    pie.legend.position = "r"
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.dataLabels.showLeaderLines = True
    pie.height = 12
    pie.width = 16
    ws.add_chart(pie, "E2")


def write_realism_sheet(ws, latest, raw_rows, derived_rows, bin_headers, hypso_headers):
    ws["A1"] = "Slope Histogram"
    ws["D1"] = "Hypsometric Curve"
    ws["G1"] = "Realism Trend"
    ws["A2"] = "Slope angle (degrees)"
    ws["B2"] = "Vertex count (cells)"
    for degree, header in enumerate(bin_headers, start=3):
        ws.cell(degree, 1, degree - 3)
        ws.cell(degree, 2, latest.get(header, 0) or 0)
    ws["D2"] = "Normalized height threshold (%)"
    ws["E2"] = "Area above threshold (%)"
    for i, header in enumerate(hypso_headers, start=3):
        ws.cell(i, 4, i - 3)
        ws.cell(i, 5, latest.get(header, 0) or 0)
    realism_headers = ["Run", "Moved volume (height units)", "Mean slope (deg)", "P95 slope (deg)", "Hypsometric AUC (%)"]
    realism_rows = [[r.get("Run"), r.get("TotalVolumeMoved"), r.get("MeanSlope_deg"), r.get("P95Slope_deg"), r.get("HypsometricAUC")] for r in derived_rows]
    append_table(ws, realism_headers, realism_rows, start_row=2, start_col=7)
    for row in [1, 2]:
        style_header(ws, row)
    slope_chart = BarChart()
    slope_chart.title = "Latest Run: Slope Histogram"
    slope_chart.add_data(Reference(ws, min_col=2, min_row=2, max_row=92), titles_from_data=True)
    slope_chart.set_categories(Reference(ws, min_col=1, min_row=3, max_row=92))
    configure_chart(
        slope_chart,
        x_title="Slope angle (degrees)",
        y_title="Vertex count (cells)",
        x_format="0",
        y_format="0",
        y_major_unit=nice_major_unit([latest.get(header, 0) or 0 for header in bin_headers]),
    )
    slope_chart.height = 10
    slope_chart.width = 18
    ws.add_chart(slope_chart, "M2")
    hypso_chart = LineChart()
    hypso_chart.title = "Latest Run: Hypsometric Curve"
    hypso_chart.add_data(Reference(ws, min_col=5, min_row=2, max_row=103), titles_from_data=True)
    hypso_chart.set_categories(Reference(ws, min_col=4, min_row=3, max_row=103))
    configure_chart(
        hypso_chart,
        x_title="Normalized height threshold (%)",
        y_title="Area above threshold (%)",
        x_format="0",
        y_format="0",
        x_major_unit=10,
        y_major_unit=10,
    )
    hypso_chart.height = 10
    hypso_chart.width = 18
    ws.add_chart(hypso_chart, "M20")
    add_line_chart(
        ws,
        "Realism Metrics by Run",
        Reference(ws, min_col=8, max_col=11, min_row=2, max_row=2 + len(realism_rows)),
        Reference(ws, min_col=7, min_row=3, max_row=2 + len(realism_rows)),
        "M38",
        x_title="Run number",
        y_title="Metric value (mixed units)",
        x_format="0",
        y_format="0.00",
    )


def write_tradeoff_sheet(ws, rows):
    sorted_rows = sorted(rows, key=lambda r: (r.get("Iterations") or 0, r.get("PhysicsTime_ms") or 0))
    headers = ["Droplet iterations (count)", "Physics simulation time (ms)", "Total runtime incl. generation (ms)", "Moved volume (height units)", "Hypsometric AUC (%)", "Mean slope (deg)", "P95 slope (deg)"]
    source_keys = ["Iterations", "PhysicsTime_ms", "TotalTime_ms", "TotalVolumeMoved", "HypsometricAUC", "MeanSlope_deg", "P95Slope_deg"]
    table = [[r.get(h) for h in source_keys] for r in sorted_rows]
    start, end = append_table(ws, headers, table)
    if table:
        iteration_values = [row[0] for row in table]
        physics_values = [row[1] for row in table]
        volume_values = [row[3] for row in table]
        add_scatter_chart(
            ws,
            "Elbow Curve: Iterations vs Volume Moved",
            Reference(ws, min_col=1, min_row=start + 1, max_row=end),
            Reference(ws, min_col=4, min_row=start + 1, max_row=end),
            "H2",
            "Droplet iterations (count)",
            "Moved volume (height units)",
            x_format="0",
            y_format="0.00",
            x_major_unit=nice_major_unit(iteration_values),
            y_major_unit=nice_major_unit(volume_values),
            series_title="Moved volume (height units)",
        )
        add_scatter_chart(
            ws,
            "Cost vs Realism",
            Reference(ws, min_col=2, min_row=start + 1, max_row=end),
            Reference(ws, min_col=4, min_row=start + 1, max_row=end),
            "H20",
            "Physics Simulation Time (ms)",
            "Moved volume (height units)",
            x_format="0",
            y_format="0.00",
            x_major_unit=nice_major_unit(physics_values),
            y_major_unit=nice_major_unit(volume_values),
            series_title="Moved volume (height units)",
        )


def write_correlations_sheet(ws, rows):
    metric_names = [
        "ExperimentalValue",
        "GridSize",
        "Iterations",
        "GenTime_ms",
        "ErosionTime_ms",
        "ThermalTime_ms",
        "PhysicsTime_ms",
        "TotalTime_ms",
        "MemoryEstimate_MB",
        "TotalVolumeMoved",
        "MeanSlope_deg",
        "P95Slope_deg",
        "HypsometricAUC",
    ]
    ws.append(["Metric", *metric_names])
    for metric_a in metric_names:
        ws.append([
            metric_a,
            *[
                correlation([r.get(metric_a) for r in rows], [r.get(metric_b) for r in rows])
                for metric_b in metric_names
            ],
        ])
    style_header(ws)
    for row in ws.iter_rows(min_row=2, min_col=2):
        for cell in row:
            cell.number_format = "0.00"


if __name__ == "__main__":
    main()
